"""Build a business-meeting Word file: dashboards and their data sources only."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path(__file__).resolve().parent / "DigiLog-Dashboard-Data-Sources.docx"

NAVY = RGBColor(0x1E, 0x3A, 0x5F)
TEAL = RGBColor(0x0F, 0x76, 0x6E)
GRAY = RGBColor(0x47, 0x55, 0x69)
HEADER_BG = "1E3A5F"
ALT_BG = "F8FAFC"


def set_run_font(run, name="Calibri", size=11, bold=False, italic=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def shade_cell(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_borders(cell, color="CBD5E1", sz="4"):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), sz)
        el.set(qn("w:color"), color)
        tcBorders.append(el)
    tcPr.append(tcBorders)


def set_cell_width(cell, cm):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcW = OxmlElement("w:tcW")
    tcW.set(qn("w:w"), str(int(cm * 567)))
    tcW.set(qn("w:type"), "dxa")
    tcPr.append(tcW)


def add_field(paragraph, instr):
    run = paragraph.add_run()
    r = run._r
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = instr
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    r.append(begin)
    r.append(text)
    r.append(sep)
    r.append(end)


def space(p, before=0, after=8):
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p.paragraph_format.line_spacing = 1.08


def h(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(16 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(8 if level == 1 else 6)
    if level == 1:
        p.paragraph_format.page_break_before = True
    for run in p.runs:
        run.font.color.rgb = NAVY
        run.font.name = "Calibri"
        run.font.size = Pt(18 if level == 1 else 13)
        run.bold = True
    return p


def para(doc, text, *, italic=False, size=11, after=8):
    p = doc.add_paragraph()
    space(p, 0, after)
    run = p.add_run(text)
    set_run_font(run, size=size, italic=italic, color=GRAY if italic else None)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.clear()
    space(p, 1, 3)
    run = p.add_run(text)
    set_run_font(run, size=11)
    return p


HEADERS = ["Dashboard view", "Data source", "How this source is used on the dashboard"]
COL_W = [6.2, 7.2, 9.8]


def add_table(doc, rows):
    table = doc.add_table(rows=1 + len(rows), cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else OxmlElement("w:tblPr")
    tblW = OxmlElement("w:tblW")
    tblW.set(qn("w:w"), "13200")
    tblW.set(qn("w:type"), "dxa")
    tblPr.append(tblW)

    for j, head in enumerate(HEADERS):
        cell = table.cell(0, j)
        cell.text = ""
        p = cell.paragraphs[0]
        space(p, 3, 3)
        run = p.add_run(head)
        set_run_font(run, size=9, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        shade_cell(cell, HEADER_BG)
        set_cell_borders(cell)
        set_cell_width(cell, COL_W[j])
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = table.cell(i + 1, j)
            cell.text = ""
            p = cell.paragraphs[0]
            space(p, 3, 3)
            run = p.add_run(str(val))
            set_run_font(run, size=9)
            if i % 2 == 1:
                shade_cell(cell, ALT_BG)
            set_cell_borders(cell)
            set_cell_width(cell, COL_W[j])
    spacer = doc.add_paragraph()
    space(spacer, 0, 10)


def configure(doc):
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    sec = doc.sections[0]
    sec.page_width = Cm(29.7)
    sec.page_height = Cm(21.0)
    sec.top_margin = Cm(1.6)
    sec.bottom_margin = Cm(1.6)
    sec.left_margin = Cm(1.6)
    sec.right_margin = Cm(1.6)
    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = hp.add_run("DigiLog v2  |  Dashboard Data Sources  |  Internal")
    set_run_font(run, size=9, color=GRAY)
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run("Zuari Industries  |  GSMA  |  Page ")
    set_run_font(run, size=9, color=GRAY)
    add_field(fp, "PAGE")
    run = fp.add_run(" of ")
    set_run_font(run, size=9, color=GRAY)
    add_field(fp, "NUMPAGES")


def title_page(doc):
    for _ in range(2):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("ZUARI INDUSTRIES  |  GSMA")
    set_run_font(run, size=12, bold=True, color=TEAL)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    space(p, 14, 4)
    run = p.add_run("DigiLog")
    set_run_font(run, size=36, bold=True, color=NAVY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    space(p, 0, 8)
    run = p.add_run("Dashboard Data Sources")
    set_run_font(run, size=22, color=NAVY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(
        "Where each dashboard gets its numbers: Forms Hub, Data Upload, or the MSSQL database."
    )
    set_run_font(run, size=12, italic=True, color=GRAY)

    meta = [
        ("Prepared for", "Business review"),
        ("Product", "DigiLog v2"),
        ("Date", "24 October 2026"),
        ("Coverage", "All eight BI dashboards"),
        ("Classification", "Internal"),
    ]
    t = doc.add_table(rows=len(meta), cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (k, v) in enumerate(meta):
        t.cell(i, 0).text = ""
        t.cell(i, 1).text = ""
        pk = t.cell(i, 0).paragraphs[0]
        pv = t.cell(i, 1).paragraphs[0]
        pk.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = pk.add_run(k)
        set_run_font(run, size=11, color=GRAY)
        run = pv.add_run(v)
        set_run_font(run, size=11, bold=True, color=NAVY)
        shade_cell(t.cell(i, 0), "F1F5F9")
        set_cell_borders(t.cell(i, 0), "E2E8F0")
        set_cell_borders(t.cell(i, 1), "E2E8F0")

    p = doc.add_paragraph()
    space(p, 18, 4)
    run = p.add_run("How to read this document")
    set_run_font(run, size=13, bold=True, color=NAVY)
    para(
        doc,
        "Each chapter is one dashboard. Data source is only one of three places:",
    )
    bullet(doc, "Forms Hub  -  written as Form > Logbook, for example Operations Logbook > Lab Logbook")
    bullet(doc, "Data Upload  -  written as Data Upload > section > file slot")
    bullet(doc, "MSSQL database  -  Cane Operations only (AIRACANE)")
    para(
        doc,
        "The last column explains how that source appears on the dashboard.",
    )


def toc(doc):
    h(doc, "Contents")
    para(
        doc,
        "In Word: click the field below, press F9 (or right-click, Update Field) to fill page numbers.",
        italic=True,
        size=10,
    )
    p = doc.add_paragraph()
    add_field(p, r'TOC \o "1-2" \h \z \u')
    for item in [
        "1. Management Overview",
        "2. Cane Operations",
        "3. Cane Quality (Brix)",
        "4. Centre Maturity",
        "5. Milling Operations",
        "6. Power Plant Operations",
        "7. Distillery Operations",
        "8. Grower Intelligence",
    ]:
        bullet(doc, item)


def body(doc):
    h(doc, "1. Management Overview")
    para(
        doc,
        "Plant scorecard in five rows: Cane, Milling, Sugar, Power, Distillery. "
        "Each cell is a period total or average, compared with the selected prior window.",
    )
    add_table(
        doc,
        [
            [
                "Cane Indent (Q), Cane Purchase (Q), Overrun Gate / Overrun Center",
                "Data Upload > Management Dashboard > Centre Indent & Purchase",
                "Indent and purchase quantities by centre and by gate. Overrun is purchase against indent for each side.",
            ],
            [
                "Yard Bal. (8AM)",
                "Operations Logbook > Lab Logbook",
                "Cane lying in the factory yard at 8 AM. Shown on the Cane row of the scorecard.",
            ],
            [
                "Pol in Cane %",
                "DS Logbook > Lab Logbook",
                "Sucrose in cane as received for crushing. Cane-row KPI and trend.",
            ],
            [
                "Milling and Sugar row (crush, maceration, mixed juice, DMF, bagasse, DS/RS cane, sugar, recovery, filter cake)",
                "Data Upload > Management Dashboard > DMR workbook",
                "Daily mill and sugar manufacturing figures for the scorecard and daily trends. Lab DS, RS and Operations logbooks fill a reading when the DMR workbook does not have it.",
            ],
            [
                "Power/Cane, Steam/Cane, Power/Sugar, Steam/Sugar, and the Power row",
                "Power Details > Power Logbook  and  Steam Details > Power Logbook",
                "Generation, export, in-house use, steam raised, steam to sugar, steam per bagasse, and specific steam. Crush and sugar from the DMR workbook are used only to express power and steam per quintal.",
            ],
            [
                "Distillery row (syrup and molasses used, ethanol, recovery, stock, efficiency, TRS / FS)",
                "Distillery Operations Form > Distillery Operations",
                "Same daily distillery figures as the Distillery dashboard, summarised on the executive scorecard.",
            ],
        ]
    )

    h(doc, "2. Cane Operations")
    para(
        doc,
        "Gate purchase, parchie, transport, overrun, yard holding, and cut-to-crush. "
        "This dashboard reads only from the MSSQL database (AIRACANE).",
    )
    add_table(
        doc,
        [
            [
                "Procurement Summary",
                "MSSQL database (AIRACANE)",
                "Cut-to-crush process map: farm/gate vehicles and cane, centre vehicles and cane, trips, holding, transit, yard wait, mill unloading, and vehicles over standard hold time.",
            ],
            [
                "Gate 1",
                "MSSQL database (AIRACANE)",
                "Cane purchased at gate, number of purchy, average purchy size, parchie overrun by cart/trolly/truck, purchase split and daily purchase trend.",
            ],
            [
                "Gate 2",
                "MSSQL database (AIRACANE)",
                "Yard holding by transport mode: vehicle count, min/avg/max holding, and vehicles exceeding 8 hours.",
            ],
            [
                "Center Purchase",
                "MSSQL database (AIRACANE)",
                "Centre cane purchased, purchy count, average parchi size, trips, overrun by mode, top/bottom centres, and daily trend.",
            ],
            [
                "Vehicle Handling",
                "MSSQL database (AIRACANE)",
                "Vehicles handled at centres, mode-wise split, handling trend, and top/bottom centre tables.",
            ],
            [
                "Vehicle Holding and Vehicle Holding 2",
                "MSSQL database (AIRACANE)",
                "Average holding hours by mode, holding trend, and centre-wise holding tables.",
            ],
            [
                "Truck Transit and Truck Holding",
                "MSSQL database (AIRACANE)",
                "Truck transit hours, truck holding hours, and trip count by centre. Truck Holding can be filtered by challan.",
            ],
        ]
    )

    h(doc, "3. Cane Quality (Brix)")
    para(doc, "Field stalk Brix and yard consignment quality. Two tabs: Field Sampling and Yard Sampling.")
    add_table(
        doc,
        [
            [
                "Field Sampling",
                "GSMA Field Brix Sampling > Brix Sampling Forms",
                "Sample count, average top / mid / bottom Brix, maturity, crop standing, soil, variety, land type, and hydration (waterlogged vs no water).",
            ],
            [
                "Yard Sampling",
                "GSMA Yard Brix Sampling > Brix Sampling Forms",
                "Yard sample count, average mid Brix, daily trend, Brix by carrier, and consignment condition. Delivery point is the yard filter.",
            ],
        ]
    )

    h(doc, "4. Centre Maturity")
    para(
        doc,
        "Centre-wise indent vs purchase for the selected window, compared with the prior window.",
    )
    add_table(
        doc,
        [
            [
                "Indent Qty, Purchase Qty, centre ranking and maturity",
                "Data Upload > Management Dashboard > Centre Indent & Purchase",
                "Headline indent and purchase. Each centre row is purchase against indent (maturity). Centres are ranked and coloured from that ratio.",
            ],
        ]
    )

    h(doc, "5. Milling Operations")
    para(
        doc,
        "Mill stoppages and mill-house temperatures. Four areas: Mill Outage, Equipment Temperature, Shredder and OTG Temp, Lube & Roller Temp.",
    )
    add_table(
        doc,
        [
            [
                "Mill Outage",
                "Mill Stoppages > Mill Logbook",
                "Stoppage hours, split by section and machinery, plus the event list.",
            ],
            [
                "Equipment Temperature",
                "Equipment Temperature > Mill Logbook",
                "Daily equipment temperatures (motors, gearboxes, bearings). Equipment display names come from Data Upload > Milling Operations > Mill Equipment Mapping.",
            ],
            [
                "Shredder and OTG Temp",
                "Shredder and OTG > Mill Logbook",
                "Shredder and OTG bearing temperatures. Names come from Data Upload > Milling Operations > Shredder Equipment Mapping.",
            ],
            [
                "Lube & Roller Temp",
                "Lube Pressure and Roller Temp > Mill Logbook",
                "Lube pressure and roller temperatures. Names come from Data Upload > Milling Operations > Lube & Roller Mapping.",
            ],
        ]
    )

    h(doc, "6. Power Plant Operations")
    para(
        doc,
        "Generation, consumption, steam, outage and process from Power Logbook forms.",
    )
    add_table(
        doc,
        [
            [
                "Power Summary, Generation, Consumption",
                "Power Details > Power Logbook",
                "Power generated by each TG set, export to grid/sugar/distillery, in-house use, import, and crush.",
            ],
            [
                "Steam Summary and Steam Consumption",
                "Steam Details > Power Logbook",
                "Steam raised on 150 / 70 / 35 TPH boilers, steam to sugar, bagasse fired, and steam used by each TG set.",
            ],
            [
                "Outage",
                "Stoppage Details > Power Logbook",
                "Stoppage hours by section, sub-section, machinery and category. Filterable outage list.",
            ],
            [
                "Process",
                "Power Details, Steam Details and Stoppage Details > Power Logbook",
                "Combined day-wise view of power, steam and outages.",
            ],
        ]
    )

    h(doc, "7. Distillery Operations")
    para(
        doc,
        "Ethanol volume, fermentation sugar, efficiency, wash, molasses stock and ethanol stock.",
    )
    add_table(
        doc,
        [
            [
                "Whole dashboard (KPI strip and charts)",
                "Distillery Operations Form > Distillery Operations",
                "Ethanol produced, syrup/molasses used, fermentation and distillation efficiency, wash distilled, B and C molasses in store, ethanol in store, TRS / FS, and recovery. Operating days are days with ethanol production.",
            ],
        ]
    )

    h(doc, "8. Grower Intelligence")
    para(
        doc,
        "Grower performance, indent vs supply, purchy dishonour, staff hierarchy and failure by date.",
    )
    add_table(
        doc,
        [
            [
                "Grower Performance, Purchy Dishonour, Dishonour Drilldown, Failure by Date",
                "Data Upload > Grower Intelligence > Grower Details Season",
                "Grower bonding, indent, weighed supply, indent failure, and dishonour quantity / %. Society and loyalty filters. Drill from society to village to grower, and failure by date.",
            ],
            [
                "Staff Drilldown",
                "Data Upload > Grower Intelligence > Staff wise Bonding",
                "Zone head to village staff hierarchy. Selecting a person filters growers and supply under that person. Grower figures still come from Grower Details Season.",
            ],
        ]
    )


def main():
    doc = Document()
    configure(doc)
    title_page(doc)
    toc(doc)
    body(doc)
    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()

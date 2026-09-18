#!/usr/bin/env python3
"""
Extract DS House mechanical cards for the 18-09-26 batch.

Sources:
  Extraction_files-sugar/new files - 180926/
    BOILING HOUSE MECHANICAL  EQUIPMENT LIST -30-07-2026.xlsx
    ANOTHER SHEET OF EQUP. HISTORY SHEET.xlsx

Scope:
  Mechanical tags ZIL/GSM/DS/117–182 (66 history cards).
  History-only ZIL/GSM/DS/149 gets a synthetic hierarchy row.

Matching (same as mill-house mechanical):
  tag + Sub Equipment. Duplicate tags stay as separate leaves.
  One history sheet maps to one leaf; do not copy onto every sibling.

Outputs (extraction folder + backlog copy):
  ds-house-mechanical-hierarchy-180926.xlsx
  ds-house-mechanical-equipment-history-180926.xlsx
"""

from __future__ import annotations

import re
import shutil
import uuid
from collections import defaultdict
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill

from equipment_history_extract_lib import cell_text, norm
from extract_mill_house_electrical import (
    HISTORY_COLUMNS,
    LIFE_COLUMNS,
    SCHEDULE_COLUMNS,
    SPEC_COLUMNS,
    extract_maintenance_history,
    extract_maintenance_schedule,
    extract_specification_rows,
    insert_tag_column,
    prepend_source_file,
    style_header_row,
    write_data_sheet,
)

ROOT = Path(__file__).resolve().parent.parent.parent
SRC_DIR = ROOT / "Extraction_files-sugar" / "new files - 180926"
BACKLOG_DIR = ROOT / "DigiLog" / "backend" / "backlog-data" / "mill data" / "migration files-18-09-26"

HIERARCHY_FILE = SRC_DIR / "BOILING HOUSE MECHANICAL  EQUIPMENT LIST -30-07-2026.xlsx"
DATA_FILE = SRC_DIR / "ANOTHER SHEET OF EQUP. HISTORY SHEET.xlsx"

HIER_OUT_NAME = "ds-house-mechanical-hierarchy-180926.xlsx"
HIST_OUT_NAME = "ds-house-mechanical-equipment-history-180926.xlsx"

SKIP_SHEETS = {"index", "summary index", "summary link", "sheet1"}
TAG_MIN = 117
TAG_MAX = 182
DS_149 = "zil/gsm/ds/149"

HIERARCHY_HEADERS = [
    "Sr.No.",
    "Plant",
    "Section",
    "Location",
    "Main Equipment",
    " Sub Equipment",
    "Department",
    " History card Tag nas.",
    "History card Location",
]

YES_FILL = PatternFill("solid", fgColor="C6EFCE")
NO_FILL = PatternFill("solid", fgColor="FFC7CE")
YES_FONT = Font(bold=True, color="006100")
NO_FONT = Font(bold=True, color="9C0006")


def norm_tag(value: str) -> str:
    return re.sub(r"\s+", "", str(value or "")).strip().lower()


def clean_name(value) -> str:
    return re.sub(r"\s+", " ", str(value or "").replace("\n", " ")).strip()


def ds_tag_number(tag: str) -> int | None:
    match = re.search(r"zil/gsm/ds/(\d+)", norm_tag(tag))
    if not match:
        return None
    return int(match.group(1))


def in_batch(tag: str) -> bool:
    number = ds_tag_number(tag)
    return number is not None and TAG_MIN <= number <= TAG_MAX


def is_mechanical_dept(dept: str) -> bool:
    upper = str(dept or "").upper().replace(" ", "")
    return "MECHANICAL" in upper or "MECHNICAL" in upper or upper.startswith("MECH")


def split_tags(raw: str) -> list[str]:
    parts: list[str] = []
    for line in re.split(r"[\n\r;]+", str(raw or "")):
        text = line.strip()
        if text and "/" in text:
            parts.append(text)
    return parts


def pick_mechanical_tag(tags: list[str]) -> str | None:
    cleaned = [t.strip() for t in tags if t and "/" in t]
    if not cleaned:
        return None
    if len(cleaned) == 1:
        return cleaned[0]
    gsm = [t for t in cleaned if norm_tag(t).startswith("zil/gsm/")]
    if gsm:
        return gsm[0]
    non_inst = [t for t in cleaned if not norm_tag(t).startswith("zil/sug")]
    if non_inst:
        return non_inst[0]
    return cleaned[-1]


def should_skip_sheet(name: str) -> bool:
    return norm(name.strip()).lower() in SKIP_SHEETS


def row_value_after_label(ws, row_idx: int, label_col: int) -> str:
    for col in range(label_col + 1, ws.max_column + 1):
        val = cell_text(ws.cell(row_idx, col).value)
        if val:
            return val
    return ""


def equipment_name_key(value) -> str:
    s = str(value or "").lower()
    s = s.replace(".", " ").replace("-", " ").replace("_", " ")
    s = re.sub(r"\s+", " ", s).strip()
    s = re.sub(r"\bno\s+", "", s)
    s = s.replace("cardian", "carding")
    return re.sub(r"\s+", " ", s).strip()


def name_key_variants(value) -> set[str]:
    key = equipment_name_key(value)
    if not key:
        return set()
    variants = {key}
    if key.startswith("cane "):
        variants.add(key[5:])
    return variants


def names_match(a, b) -> bool:
    return bool(name_key_variants(a) & name_key_variants(b))


def trailing_number(value) -> str:
    key = equipment_name_key(value)
    match = re.search(r"(\d+)$", key)
    return match.group(1) if match else ""


def pick_hierarchy_row_for_sheet(fields: dict, sheet_name: str, metas: list[dict]) -> dict | None:
    """Match one hierarchy leaf: same tag group, same sub-equipment name."""
    card = (fields or {}).get("NAME OF EQUIPMENT") or ""
    title = sheet_name or ""
    for meta in metas:
        sub = meta.get("sub_equipment") or ""
        if names_match(card, sub) or names_match(title, sub):
            return meta
    for meta in metas:
        sub = meta.get("sub_equipment") or ""
        main = meta.get("main_equipment") or ""
        if names_match(sub, main) and (names_match(card, main) or names_match(title, main)):
            return meta
    card_num = trailing_number(card) or trailing_number(title)
    if card_num:
        numbered = [meta for meta in metas if trailing_number(meta.get("sub_equipment") or "") == card_num]
        if len(numbered) == 1:
            return numbered[0]
    return None


def extract_life_fields(ws) -> dict[str, str]:
    fields = {
        "NAME OF EQUIPMENT": ws.title.strip(),
        "LOCATION": "",
        "EQUIPMENT TAG NO": "",
        "DATE OF COMMISSIONING": "",
    }
    life_row = None
    spec_row = ws.max_row + 1

    for row_idx in range(1, min(ws.max_row, 40) + 1):
        for col in (1, 2, 3, 4, 5):
            text = norm(ws.cell(row_idx, col).value)
            if not text:
                continue
            if life_row is None and "EQUIPMENT LIFE HISTORY" in text:
                life_row = row_idx
            if "EQUIPMENT SPECIFICATION" in text:
                spec_row = row_idx
                break

    scan_start = (life_row + 1) if life_row is not None else 1
    scan_end = spec_row if spec_row <= ws.max_row else min(ws.max_row, 20)
    equip_no = ""

    for row_idx in range(scan_start, scan_end):
        label = ""
        label_col = 1
        for col in range(1, min(6, ws.max_column + 1)):
            raw = norm(ws.cell(row_idx, col).value)
            if raw:
                label = raw
                label_col = col
                break
        if not label:
            continue
        value = row_value_after_label(ws, row_idx, label_col)
        if "NAME OF EQUIPMENT" in label:
            fields["NAME OF EQUIPMENT"] = value or fields["NAME OF EQUIPMENT"]
        elif "LOCATION" in label:
            fields["LOCATION"] = value
        elif "EQUIPMENT TAG" in label or label.startswith("TAG NO") or label == "TAG NAME":
            fields["EQUIPMENT TAG NO"] = value
        elif "EQUIPMENT NO" in label and "TAG" not in label:
            equip_no = value
        elif "DATE OF COMMISSIONING" in label or "COMMISSIONING" in label:
            fields["DATE OF COMMISSIONING"] = value

    if not fields["EQUIPMENT TAG NO"] and equip_no:
        fields["EQUIPMENT TAG NO"] = equip_no
    if not fields["NAME OF EQUIPMENT"]:
        fields["NAME OF EQUIPMENT"] = ws.title.strip()
    return fields


def load_hierarchy_rows(path: Path) -> list[dict]:
    if not path.exists():
        raise FileNotFoundError(f"Hierarchy file not found: {path}")
    workbook = openpyxl.load_workbook(path, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    rows: list[dict] = []
    for row_idx in range(3, sheet.max_row + 1):
        dept = clean_name(sheet.cell(row_idx, 7).value)
        if not is_mechanical_dept(dept):
            continue
        mech_tag = pick_mechanical_tag(split_tags(str(sheet.cell(row_idx, 8).value or "")))
        if not mech_tag or not in_batch(mech_tag):
            continue
        rows.append(
            {
                "sr_no": clean_name(sheet.cell(row_idx, 1).value),
                "plant": clean_name(sheet.cell(row_idx, 2).value) or "SUGAR",
                "section": clean_name(sheet.cell(row_idx, 3).value),
                "location": clean_name(sheet.cell(row_idx, 4).value),
                "main_equipment": clean_name(sheet.cell(row_idx, 5).value),
                "sub_equipment": clean_name(sheet.cell(row_idx, 6).value),
                "department": "MECHANICAL",
                "raw_tag": mech_tag,
                "hist_location": clean_name(sheet.cell(row_idx, 9).value),
                "hierarchy_row": row_idx,
                "synthetic": False,
            }
        )
    workbook.close()
    return rows


def rows_by_tag(rows: list[dict]) -> dict[str, list[dict]]:
    grouped: dict[str, list[dict]] = defaultdict(list)
    for meta in rows:
        grouped[norm_tag(meta["raw_tag"])].append(meta)
    return grouped


def centrifugal_location(rows: list[dict]) -> tuple[str, str]:
    for meta in rows:
        if "CENTRIF" in (meta.get("main_equipment") or "").upper() and meta.get("sub_equipment", "").endswith("1"):
            return meta.get("location") or "", meta.get("main_equipment") or "A-CENTRIFUGAL MACHINE"
    for meta in rows:
        if "CENTRIF" in (meta.get("main_equipment") or "").upper():
            return meta.get("location") or "", meta.get("main_equipment") or "A-CENTRIFUGAL MACHINE"
    return "", "A-CENTRIFUGAL MACHINE"


def ensure_ds149(hierarchy_rows: list[dict], fields: dict | None) -> dict:
    for meta in hierarchy_rows:
        if norm_tag(meta["raw_tag"]) == DS_149:
            return meta
    location, main = centrifugal_location(hierarchy_rows)
    name = ""
    hist_loc = ""
    if fields:
        name = fields.get("NAME OF EQUIPMENT") or ""
        hist_loc = fields.get("LOCATION") or ""
    meta = {
        "sr_no": "149-SYN",
        "plant": "SUGAR",
        "section": "DS HOUSE",
        "location": location or "A-CENTRIFUGAL MACHINE",
        "main_equipment": main,
        "sub_equipment": name or "Batch machine 1 ton",
        "department": "MECHANICAL",
        "raw_tag": "ZIL/GSM/DS/149",
        "hist_location": hist_loc,
        "hierarchy_row": 0,
        "synthetic": True,
    }
    hierarchy_rows.append(meta)
    return meta


def save_workbook_to(out: openpyxl.Workbook, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    fallback = path.with_name(f"{path.stem}-updated{path.suffix}")
    for candidate in (path, fallback):
        try:
            out.save(candidate)
            return candidate
        except PermissionError:
            continue
    raise PermissionError(
        f"Could not write {path.name} or {fallback.name}. Close them in Excel and retry."
    )


def copy_to_backlog(src: Path) -> Path:
    BACKLOG_DIR.mkdir(parents=True, exist_ok=True)
    dest = BACKLOG_DIR / src.name
    shutil.copy2(src, dest)
    return dest


def write_hierarchy_xlsx(hierarchy_rows: list[dict], dest: Path) -> Path:
    out = openpyxl.Workbook()
    ws = out.active
    ws.title = "Sheet1"
    ws.append(HIERARCHY_HEADERS)
    for meta in hierarchy_rows:
        ws.append(
            [
                meta["sr_no"],
                meta["plant"],
                meta["section"],
                meta["location"],
                meta["main_equipment"],
                meta["sub_equipment"],
                meta["department"],
                meta["raw_tag"],
                meta["hist_location"],
            ]
        )
    style_header_row(ws)
    ws.freeze_panes = "A2"
    return save_workbook_to(out, dest)


def main() -> None:
    print(f"Hierarchy: {HIERARCHY_FILE}")
    print(f"Data:      {DATA_FILE}")
    if not DATA_FILE.exists():
        raise FileNotFoundError(f"Data file not found: {DATA_FILE}")

    hierarchy_rows = load_hierarchy_rows(HIERARCHY_FILE)
    print(f"Hierarchy rows in batch (before DS/149): {len(hierarchy_rows)}")

    src = openpyxl.load_workbook(DATA_FILE, data_only=True)
    source_name = DATA_FILE.name
    ds149_fields = None
    for sheet_name in src.sheetnames:
        if should_skip_sheet(sheet_name):
            continue
        fields = extract_life_fields(src[sheet_name])
        if norm_tag(fields.get("EQUIPMENT TAG NO")) == DS_149:
            ds149_fields = fields
            break
    ensure_ds149(hierarchy_rows, ds149_fields)
    hierarchy_by_tag = rows_by_tag(hierarchy_rows)
    print(f"Hierarchy rows after synthetic DS/149: {len(hierarchy_rows)}")

    all_mappings: list[tuple[str, str, str, str, str]] = []
    all_life: list[list[str]] = []
    all_spec: list[list[str]] = []
    all_schedule: list[list[str]] = []
    all_history: list[list[str]] = []
    matched_sheets: dict[str, list[str]] = {}
    matched_by_row: dict[int, str] = {}
    unmatched_sheets: list[str] = []
    unmatched_names: list[str] = []

    for sheet_name in src.sheetnames:
        if should_skip_sheet(sheet_name):
            continue

        ws = src[sheet_name]
        fields = extract_life_fields(ws)
        equip_tag = fields["EQUIPMENT TAG NO"]
        nt = norm_tag(equip_tag)
        if not in_batch(equip_tag):
            unmatched_sheets.append(f"{sheet_name} (tag={equip_tag or '(blank)'} out of batch)")
            continue

        metas = hierarchy_by_tag.get(nt) or []
        if not metas:
            unmatched_sheets.append(f"{sheet_name} (tag={equip_tag or '(blank)'})")
            continue

        matched_sheets.setdefault(nt, [])
        if sheet_name in matched_sheets[nt]:
            continue
        matched_sheets[nt].append(sheet_name)

        meta = pick_hierarchy_row_for_sheet(fields, sheet_name, metas)
        if len(metas) == 1:
            meta = metas[0]
        chosen = meta or metas[0]
        out_tag = chosen["raw_tag"]
        # Unique tag: always use hierarchy Sub Equipment so import combo key matches.
        # Duplicate tag without a name hit: keep sheet name (mill-house behaviour).
        if meta:
            sub_equipment = chosen["sub_equipment"]
            matched_by_row[meta["hierarchy_row"]] = sheet_name
        else:
            sub_equipment = fields["NAME OF EQUIPMENT"] or chosen["sub_equipment"]
            unmatched_names.append(
                f"{sheet_name} tag={out_tag} name={fields['NAME OF EQUIPMENT']}"
            )

        sheet_id = uuid.uuid4().hex[:12]
        all_mappings.append((source_name, sheet_name, sheet_id, out_tag, sub_equipment))
        location = fields["LOCATION"]
        if not location:
            location = chosen.get("hist_location") or chosen.get("location") or ""
        all_life.append(
            [
                source_name,
                sheet_id,
                sheet_name,
                out_tag,
                fields["NAME OF EQUIPMENT"] or sub_equipment,
                location,
                fields["DATE OF COMMISSIONING"],
                chosen["main_equipment"],
                sub_equipment,
                chosen["department"],
                chosen.get("hist_location") or "",
            ]
        )
        all_spec.extend(prepend_source_file(
            insert_tag_column(
                extract_specification_rows(ws, sheet_id, sheet_name),
                out_tag,
            ),
            source_name,
        ))
        all_schedule.extend(prepend_source_file(
            insert_tag_column(
                extract_maintenance_schedule(ws, sheet_id, sheet_name),
                out_tag,
            ),
            source_name,
        ))
        all_history.extend(prepend_source_file(
            insert_tag_column(
                extract_maintenance_history(ws, sheet_id, sheet_name),
                out_tag,
            ),
            source_name,
        ))

    src.close()

    print(f"History sheets matched by tag: {len(matched_sheets)}")
    print(f"Extracted cards (one per sheet): {len(all_mappings)}")
    print(f"Hierarchy rows matched by tag + sub-equipment: {len(matched_by_row)}")
    if unmatched_names:
        print("Sheets whose name did not match a sub-equipment (still extracted):")
        for line in unmatched_names:
            print(f"  - {line}")
    if unmatched_sheets:
        print("Data sheets skipped:")
        for line in unmatched_sheets:
            print(f"  - {line}")

    hier_path = write_hierarchy_xlsx(hierarchy_rows, SRC_DIR / HIER_OUT_NAME)

    hist_out = openpyxl.Workbook()
    hist_out.remove(hist_out.active)
    mech_life_columns = [*LIFE_COLUMNS, "History card Location"]

    map_ws = hist_out.create_sheet("Sheet Map")
    map_ws.append(["source file", "sheet name", "id", "EQUIPMENT TAG NO", "Sub Equipment"])
    for source_file_name, sheet_name, sheet_id, equip_tag, sub_equipment in all_mappings:
        map_ws.append([source_file_name, sheet_name, sheet_id, equip_tag, sub_equipment])
    style_header_row(map_ws)
    map_ws.freeze_panes = "A2"
    map_ws.column_dimensions["A"].width = 48
    map_ws.column_dimensions["B"].width = 36
    map_ws.column_dimensions["C"].width = 16
    map_ws.column_dimensions["D"].width = 28
    map_ws.column_dimensions["E"].width = 42

    write_data_sheet(
        hist_out,
        "EQUIPMENT LIFE HISTORY CARD",
        mech_life_columns,
        all_life,
        {
            "A": 48, "B": 16, "C": 36, "D": 28, "E": 36, "F": 22,
            "G": 20, "H": 28, "I": 36, "J": 22, "K": 28,
        },
    )
    write_data_sheet(
        hist_out,
        "EQUIPMENT SPECIFICATION",
        SPEC_COLUMNS,
        all_spec,
        {"A": 48, "B": 16, "C": 36, "D": 28, "E": 14, "F": 18, "G": 28, "H": 36},
    )
    write_data_sheet(
        hist_out,
        "MAINTENANCE SCHEDULE",
        SCHEDULE_COLUMNS,
        all_schedule,
        {
            "A": 48, "B": 16, "C": 36, "D": 28, "E": 8, "F": 22, "G": 42,
            "H": 8, "I": 8, "J": 8, "K": 10, "L": 12, "M": 8,
            "N": 10, "O": 10, "P": 10, "Q": 16,
        },
    )
    write_data_sheet(
        hist_out,
        "EQUIPMENT MAINTENANCE HISTORY",
        HISTORY_COLUMNS,
        all_history,
        {
            "A": 48, "B": 16, "C": 36, "D": 28, "E": 18, "F": 14, "G": 16,
            "H": 16, "I": 36, "J": 36, "K": 16, "L": 22, "M": 28, "N": 24,
        },
    )
    hist_path = save_workbook_to(hist_out, SRC_DIR / HIST_OUT_NAME)

    hier_copy = copy_to_backlog(hier_path)
    hist_copy = copy_to_backlog(hist_path)

    print(f"\nHIERARCHY: {hier_path} ({len(hierarchy_rows)} rows)")
    print(f"HISTORY:   {hist_path} ({len(all_mappings)} cards)")
    print(f"BACKLOG:   {hier_copy}")
    print(f"           {hist_copy}")
    print(
        f"Sheet Map ({len(all_mappings)}), LIFE ({len(all_life)}), "
        f"SPEC ({len(all_spec)}), SCHEDULE ({len(all_schedule)}), "
        f"HISTORY ({len(all_history)})"
    )


if __name__ == "__main__":
    main()
